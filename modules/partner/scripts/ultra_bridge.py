#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Мост между июльским импортом Ultra (коды GOG*) и синхронизацией по API (ULT*).

    python3 modules/partner/scripts/ultra_bridge.py            # разбор, без записи
    python3 modules/partner/scripts/ultra_bridge.py --apply    # проставить связь

Зачем. В июле 2026 каталог Ultra загрузили файлом `ULTRA.md (1).xlsx`
(Set_data_import/7 и /8). В колонке «Артикул» там стоял НАШ внутренний код
`GOG*`, артикула поставщика не было вовсе — так в номенклатуру попали ~22 000
карточек, которые ничем не связаны с поставщиком.

Теперь тот же каталог приходит по партнёрскому API с артикулами `ULT*` и их
uuid. Общего ключа между `GOG` и `ULT` в базе нет:

    их uuid -> TMS_MPT_IMPSRC.SRC_PID      0 совпадений
    ULT*    -> TMS_UNIVERS.CODVECHI        0 совпадений
    точное название                      252 совпадения

Если публиковать срез как есть, конвейер заведёт все 34 437 позиций заново и
задвоит ~16 000 существующих.

Ключ нашёлся в адресе картинки — он есть и там, и там, только хост разный:

    июль : cdn.ultra.md/images/webp/products/<uuid>/images/26
    API  : cdn-ultra.esempla.com/storage/webp/<uuid>.webp

Пересечение — 15 952 позиции, у всех есть карточка в номенклатуре.

ВАЖНО: uuid ТОВАРА (`BIRO26_GOODS.GUID`) для этого не годится — он совпадает
всего в 1 704 случаях. Ключ именно от КАРТИНКИ.

Что пишется:

    TMS_MPT_IMPSRC.SRC_PID         <- uuid товара в Ultra
    TMS_MPT_IMPSRC.SRC_ARTICOL     <- артикул ULT*
    TMS_MPT_IMPSRC.SRC_SOURCE_CODE <- 'ULTRA'
    MATCH_STATUS                   <- 'IMG_UUID'

`TMS_UNIVERS.CODVECHI` не трогаем: код `GOG*` показан на витрине, стоит в
адресах товара и в истории продаж. Цены тут тоже не пишутся — это работа
штатного конвейера импорта, который после моста наконец узнает товар.
"""
from __future__ import annotations

import argparse
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, ROOT)

UUID_RE = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
JULY_XLSX = "/Users/pt/Projects.AI/BIRO26/Set_data_import/8/ULTRA.md (1).xlsx"


def july_map(path: str) -> dict[str, str]:
    """uuid картинки -> код GOG из июльской выгрузки."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, str] = {}
    for sn in wb.sheetnames:
        it = wb[sn].iter_rows(values_only=True)
        next(it, None)
        for r in it:
            if not r or not r[1]:
                continue
            m = UUID_RE.search(str(r[0] or ""))
            if m:
                out.setdefault(m.group(), str(r[1]).strip())
    wb.close()
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=JULY_XLSX)
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    from dotenv import load_dotenv
    load_dotenv(os.path.join(ROOT, ".env"))
    from models.biro26_db import Biro26DB

    old = july_map(args.xlsx)
    print(f"июльская выгрузка : {len(old)} картинок с uuid")

    db = Biro26DB()
    r = db.execute_query(
        "SELECT PHOTO_URL, ARTICOL, GUID FROM BIRO26_GOODS WHERE SHEET = 'ULTRA'",
        timeout=300)
    if not r.get("success"):
        sys.exit(f"чтение буфера: {r.get('message')}")
    api = {}
    for photo, art, guid in r["data"]:
        m = UUID_RE.search(photo or "")
        if m:
            api.setdefault(m.group(), (art, guid))
    print(f"срез по API       : {len(api)} товаров с uuid картинки")

    r = db.execute_query(
        "SELECT CODVECHI, COD FROM TMS_UNIVERS WHERE TIP = 'P' AND CODVECHI LIKE 'GOG%'",
        timeout=300)
    gog2cod = {a: b for a, b in r["data"]}

    pairs = []
    for uu in set(old) & set(api):
        cod = gog2cod.get(old[uu])
        if cod:
            art, guid = api[uu]
            pairs.append((cod, art, guid))
    print(f"мост GOG <-> ULT  : {len(pairs)} карточек")

    if not args.apply:
        print("\nэто был разбор — записи не делались (нужен --apply)")
        return

    written = 0
    for i in range(0, len(pairs), 40):
        chunk = pairs[i:i + 40]
        stmts, params = [], {}
        for j, (cod, art, guid) in enumerate(chunk):
            params[f"cod{j}"], params[f"art{j}"], params[f"pid{j}"] = cod, art, guid
            stmts.append(
                "MERGE INTO TMS_MPT_IMPSRC t USING "
                f"(SELECT :cod{j} COD FROM dual) s ON (t.COD = s.COD) "
                "WHEN MATCHED THEN UPDATE SET "
                f"t.SRC_ARTICOL = :art{j}, t.SRC_PID = :pid{j}, "
                "t.SRC_SOURCE_CODE = 'ULTRA', t.MATCH_STATUS = 'IMG_UUID', "
                "t.UPDATED_AT = SYSDATE "
                "WHEN NOT MATCHED THEN INSERT "
                "(COD, SRC_SOURCE_CODE, SRC_ARTICOL, SRC_PID, MATCH_STATUS, UPDATED_AT) "
                f"VALUES (:cod{j}, 'ULTRA', :art{j}, :pid{j}, 'IMG_UUID', SYSDATE);")
        res = db.execute_dml("BEGIN " + " ".join(stmts) + " END;", params, timeout=180)
        if not res.get("success"):
            sys.exit(f"запись блока {i}: {str(res.get('message'))[:200]}")
        written += len(chunk)
        if written % 2000 == 0 or written == len(pairs):
            print(f"  записано {written}/{len(pairs)}")
    print(f"\nготово: связано {written} карточек")


if __name__ == "__main__":
    main()
