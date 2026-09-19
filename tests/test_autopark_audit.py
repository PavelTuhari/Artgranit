"""Аудиторский контур модуля «Автопарк».

Проверяется чистая часть: `audit.py` (процедуры, оценка контролей,
находки), `audit_data.demo_population` (генератор популяции),
`excel_bi`/`audit_excel` (сборка книги). Ни Oracle, ни LibreOffice
тестам не нужны: движок не импортирует БД, а книга проверяется как файл.

Главный тест здесь — `test_selfcheck_*`: набор генерируется с известным
числом подложенных дефектов, проверка гоняется вслепую, и совпадение
«подложено» = «найдено» по каждой процедуре доказывает, что методика
ловит то, что должна, и не выдумывает лишнего.

Запуск: venv/bin/python -m pytest tests/test_autopark_audit.py -q
"""
import io
import os
import sys
import zipfile
from datetime import date

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from modules.autopark import audit, audit_data  # noqa: E402


@pytest.fixture(scope="module")
def population():
    return audit_data.demo_population()


@pytest.fixture(scope="module")
def report(population):
    return audit.run_audit(population)


# ── Изоляция модуля (правило №1 проекта) ───────────────────────────────

AUDIT_FILES = ("audit.py", "audit_data.py", "audit_excel.py", "excel_bi.py",
               "audit_controller.py", "audit_routes.py")


def _read(*parts):
    with open(os.path.join(ROOT, *parts), encoding="utf-8") as fh:
        return fh.read()


def test_shared_app_does_not_know_about_the_audit():
    src = _read("app.py")
    for marker in ("audit_routes", "AuditController", "autopark.audit",
                   "audit/report.xlsx"):
        assert marker not in src, f"общий app.py упоминает {marker}"


def test_shared_deploy_script_is_untouched():
    src = _read("deploy_oracle_objects.py")
    assert "audit" not in src.lower().replace("audit_", "")


def test_audit_engine_imports_no_database():
    """Движок обязан оставаться чистым: иначе его не прогнать без wallet."""
    src = _read("modules", "autopark", "audit.py")
    for forbidden in ("models.database", "DatabaseModel", "oracledb",
                      "from modules.autopark.store", "cx_Oracle"):
        assert forbidden not in src, f"audit.py импортирует {forbidden}"


def test_excel_toolkit_knows_nothing_about_the_business():
    src = _read("modules", "autopark", "excel_bi.py")
    for forbidden in ("рейс", "резервуар", "бензовоз", "FLT_"):
        assert forbidden not in src.replace(
            "не знает ни про рейсы", ""), f"excel_bi.py знает про {forbidden}"


# ── Самопроверка методики ──────────────────────────────────────────────

def test_selfcheck_every_injected_defect_is_found(population, report):
    """Ни одного пропуска и ни одного ложного срабатывания."""
    found = {t["id"]: t["exceptions"] for t in report["tests"]}
    mismatch = {tid: (population["injected"].get(tid, 0), found.get(tid, 0))
                for tid in set(population["injected"]) | set(found)
                if population["injected"].get(tid, 0) != found.get(tid, 0)}
    assert not mismatch, f"подложено ≠ найдено: {mismatch}"


def test_selfcheck_actually_had_something_to_find(population):
    """Страховка от «зелёного» теста на пустом наборе."""
    assert sum(population["injected"].values()) > 100
    assert len([v for v in population["injected"].values() if v]) >= 10


@pytest.mark.parametrize("params", [
    {"months": 6, "stations": 8, "trucks": 4, "drivers": 6},
    {"months": 12, "stations": 40, "trucks": 20, "drivers": 30},
    {"seed": 777, "stations": 12, "trucks": 6, "drivers": 9},
])
def test_selfcheck_holds_on_other_volumes(params):
    """Методика не подогнана под один набор."""
    pop = audit_data.demo_population(**params)
    rep = audit.run_audit(pop)
    found = {t["id"]: t["exceptions"] for t in rep["tests"]}
    for tid, expected in pop["injected"].items():
        assert found.get(tid, 0) == expected, (
            f"{tid}: подложено {expected}, найдено {found.get(tid, 0)} "
            f"при {params}")


def test_population_is_deterministic():
    a = audit_data.demo_population(4242, months=6, stations=8, trucks=4,
                                   drivers=6)
    b = audit_data.demo_population(4242, months=6, stations=8, trucks=4,
                                   drivers=6)
    assert a["injected"] == b["injected"]
    assert len(a["trips"]) == len(b["trips"])
    assert [t["pay_stored"] for t in a["trips"][:50]] == \
           [t["pay_stored"] for t in b["trips"][:50]]


# ── Процедуры ──────────────────────────────────────────────────────────

def _pop(**over):
    base = {
        "settings": {"rate_per_km": 3.5, "trip_bonus": 0.0,
                     "km_deviation_limit": 5.0, "fuel_deviation_limit": 5.0,
                     "loss_tolerance_l": 50.0},
        "rate_periods": [], "prices": {"DIESEL": 20.0}, "stations": [],
        "tanks": [], "trucks": [], "drivers": [], "trips": [],
        "trip_items": [],
    }
    base.update(over)
    return base


def test_payroll_uses_the_rate_of_the_trip_date():
    """Сентябрьская ставка не переписывает августовское начисление."""
    periods = [
        {"id": 1, "valid_from": date(2026, 1, 1), "valid_to": date(2026, 8, 31),
         "rate_per_km": 2.75, "trip_bonus": 600.0},
        {"id": 2, "valid_from": date(2026, 9, 1), "valid_to": None,
         "rate_per_km": 3.50, "trip_bonus": 0.0},
    ]
    pop = _pop(rate_periods=periods, drivers=[{"id": 1, "full_name": "Х"}],
               trips=[
                   {"id": 1, "trip_date": date(2026, 8, 15), "driver_id": 1,
                    "truck_id": 1, "status_code": "APPROVED",
                    "type_code": "DOMESTIC", "norm_km": 100.0,
                    "pay_stored": 100 * 2.75 + 600},
                   {"id": 2, "trip_date": date(2026, 9, 15), "driver_id": 1,
                    "truck_id": 1, "status_code": "APPROVED",
                    "type_code": "DOMESTIC", "norm_km": 100.0,
                    "pay_stored": 100 * 3.50},
               ])
    res = audit.test_payroll_recompute(pop)
    assert res["tested"] == 2
    assert res["exceptions"] == 0, res["rows"]


def test_payroll_flags_a_hand_edited_amount():
    pop = _pop(drivers=[{"id": 1, "full_name": "Х"}],
               trips=[{"id": 7, "trip_date": date(2026, 9, 1), "driver_id": 1,
                       "truck_id": 1, "status_code": "APPROVED",
                       "type_code": "IMPORT", "norm_km": 200.0,
                       "pay_stored": 900.0}])
    res = audit.test_payroll_recompute(pop)
    assert res["exceptions"] == 1
    assert res["impact_lei"] == pytest.approx(200.0)


def test_import_trip_gets_no_bonus():
    """Прямое исключение ТЗ: доплата только за внутренний рейс."""
    pop = _pop(drivers=[{"id": 1, "full_name": "Х"}],
               rate_periods=[{"id": 1, "valid_from": date(2026, 1, 1),
                              "valid_to": None, "rate_per_km": 3.5,
                              "trip_bonus": 600.0}],
               trips=[{"id": 1, "trip_date": date(2026, 9, 1), "driver_id": 1,
                       "truck_id": 1, "status_code": "APPROVED",
                       "type_code": "IMPORT", "norm_km": 100.0,
                       "pay_stored": 350.0}])
    assert audit.test_payroll_recompute(pop)["exceptions"] == 0


def test_draft_trips_stay_out_of_payroll_testing():
    pop = _pop(drivers=[{"id": 1, "full_name": "Х"}],
               trips=[{"id": 1, "trip_date": date(2026, 9, 1), "driver_id": 1,
                       "truck_id": 1, "status_code": "DRAFT",
                       "type_code": "DOMESTIC", "norm_km": 100.0,
                       "pay_stored": 999999.0}])
    assert audit.test_payroll_recompute(pop)["tested"] == 0


def test_rate_periods_overlap_is_an_exception():
    pop = _pop(rate_periods=[
        {"id": 1, "valid_from": date(2026, 1, 1), "valid_to": date(2026, 9, 10),
         "rate_per_km": 2.75},
        {"id": 2, "valid_from": date(2026, 9, 1), "valid_to": None,
         "rate_per_km": 3.50},
    ])
    res = audit.test_rate_period_integrity(pop)
    assert res["exceptions"] == 1
    assert "ересеч" in res["rows"][0][-1]


def test_open_previous_rate_period_is_an_exception():
    pop = _pop(rate_periods=[
        {"id": 1, "valid_from": date(2026, 1, 1), "valid_to": None,
         "rate_per_km": 2.75},
        {"id": 2, "valid_from": date(2026, 9, 1), "valid_to": None,
         "rate_per_km": 3.50},
    ])
    assert audit.test_rate_period_integrity(pop)["exceptions"] == 1


TANK = {"id": 1, "station_id": 1, "station_code": "CHI-01",
        "product_code": "DIESEL", "capacity_l": 30000, "max_fill_l": 28500,
        "min_stock_l": 6000, "current_l": 20000, "avg_daily_l": 3000}


def test_overfill_compares_history_with_the_tank_limit_not_free_space():
    """Исторический слив нельзя мерить сегодняшним остатком."""
    item = {"trip_id": 1, "tank_id": 1, "station_code": "CHI-01",
            "product_code": "DIESEL", "accepted_l": 12000.0}
    pop = _pop(tanks=[dict(TANK)], trip_items=[item],
               trips=[{"id": 1, "status_code": "APPROVED"}])
    # Свободно всего 8 500 л, но поставка исполненная: предел — 28 500 л.
    assert audit.test_overfill(pop)["exceptions"] == 0


def test_overfill_compares_pending_delivery_with_free_space():
    item = {"trip_id": 1, "tank_id": 1, "station_code": "CHI-01",
            "product_code": "DIESEL", "accepted_l": 12000.0}
    pop = _pop(tanks=[dict(TANK)], trip_items=[item],
               trips=[{"id": 1, "status_code": "PLANNED"}])
    res = audit.test_overfill(pop)
    assert res["exceptions"] == 1
    assert res["rows"][0][-1] == pytest.approx(3500.0)


def test_overfill_catches_delivery_above_the_tank_limit():
    item = {"trip_id": 1, "tank_id": 1, "station_code": "CHI-01",
            "product_code": "DIESEL", "accepted_l": 31000.0}
    pop = _pop(tanks=[dict(TANK)], trip_items=[item],
               trips=[{"id": 1, "status_code": "APPROVED"}])
    assert audit.test_overfill(pop)["exceptions"] == 1


def test_section_must_be_emptied_whole():
    truck = {"id": 1, "plate": "A1", "capacity_l": 10000,
             "norm_l_per_100km": 30.0,
             "sections": [{"id": 5, "seq_no": 1, "volume_l": 5000}]}
    ok = {"trip_id": 1, "section_id": 5, "planned_l": 5000.0,
          "station_code": "X", "product_code": "A95"}
    bad = {"trip_id": 2, "section_id": 5, "planned_l": 3200.0,
           "station_code": "X", "product_code": "A95"}
    assert audit.test_section_integrity(_pop(trucks=[truck],
                                             trip_items=[ok]))["exceptions"] == 0
    assert audit.test_section_integrity(_pop(trucks=[truck],
                                             trip_items=[bad]))["exceptions"] == 1


def test_master_data_catches_fill_above_capacity():
    tank = dict(TANK, max_fill_l=32000)
    res = audit.test_master_data(_pop(tanks=[tank]))
    assert res["exceptions"] == 1
    assert "ёмкости" in res["rows"][0][-1]


def test_master_data_catches_sections_not_matching_the_tanker():
    truck = {"id": 1, "plate": "A1", "capacity_l": 20000,
             "norm_l_per_100km": 30.0,
             "sections": [{"id": 1, "seq_no": 1, "volume_l": 5000}]}
    assert audit.test_master_data(_pop(trucks=[truck]))["exceptions"] == 1


def test_segregation_of_duties_catches_self_approval():
    trips = [
        {"id": 1, "trip_date": date(2026, 9, 1), "status_code": "APPROVED",
         "created_by": "logist1", "approved_by": "logist1"},
        {"id": 2, "trip_date": date(2026, 9, 1), "status_code": "APPROVED",
         "created_by": "logist1", "approved_by": "sef"},
    ]
    res = audit.test_segregation_of_duties(_pop(trips=trips))
    assert res["tested"] == 2 and res["exceptions"] == 1


def test_status_sequence_catches_a_skipped_stage():
    trips = [{"id": 1, "status_history": ["PLANNED", "LOAD_REQ", "LOADED",
                                          "IN_TRANSIT", "DELIVERED"]},
             {"id": 2, "status_history": ["PLANNED", "LOADED"]},
             {"id": 3, "status_history": ["PLANNED", "LOAD_REQ", "PLANNED"]}]
    res = audit.test_status_sequence(_pop(trips=trips))
    assert res["tested"] == 3 and res["exceptions"] == 2


def test_duplicate_waybill_is_an_exception():
    trips = [{"id": 1, "trip_date": date(2026, 9, 1), "truck_id": 1,
              "status_code": "APPROVED", "waybill_no": "AP-1"},
             {"id": 2, "trip_date": date(2026, 9, 2), "truck_id": 2,
              "status_code": "APPROVED", "waybill_no": "AP-1"}]
    res = audit.test_duplicates(_pop(trips=trips))
    assert res["exceptions"] == 1


# ── Оценка контролей ───────────────────────────────────────────────────

def test_control_with_design_defect_never_reads_as_effective(report):
    for ctl in report["controls"]:
        if ctl["design"] != "Эффективен":
            assert ctl["rating"] != audit.RATE_OK, ctl["id"]


def test_control_without_tested_objects_is_marked_not_tested():
    """Ноль отклонений из нуля объектов — не «эффективен»."""
    assert audit._effectiveness(0, 0) == audit.RATE_NA
    assert audit._effectiveness(0, 100) == audit.RATE_OK
    assert audit._effectiveness(1, 100) == audit.RATE_WARN
    assert audit._effectiveness(20, 100) == audit.RATE_BAD


def test_live_population_leaves_missing_fields_empty_not_faked():
    """Придуманный реквизит дал бы контроль, эффективный по построению."""
    src = _read("modules", "autopark", "audit_data.py")
    tail = src[src.index("def live_population"):]
    assert '"created_by": None' in tail
    assert '"status_history": []' in tail


def test_every_control_is_linked_to_a_procedure(report):
    ids = {t["id"] for t in report["tests"]}
    for ctl in report["controls"]:
        assert ctl["test_ids"], ctl["id"]
        for tid in ctl["test_ids"]:
            assert tid in ids, f"{ctl['id']} ссылается на {tid}"


def test_every_procedure_has_a_finding_text():
    for fn in audit.TESTS:
        tid = fn(_pop())["id"]
        assert tid in audit.FINDING_TEXT, f"для {tid} нет риска и рекомендации"
        assert audit.FINDING_TEXT[tid]["risk"]
        assert audit.FINDING_TEXT[tid]["rec"]


# ── Находки, значимость, заключение ────────────────────────────────────

def test_severity_comes_from_measured_values():
    assert audit._severity(7.0, 0.0) == audit.SEV_HIGH
    assert audit._severity(0.1, 80_000.0) == audit.SEV_HIGH
    assert audit._severity(2.0, 0.0) == audit.SEV_MED
    assert audit._severity(0.1, 6_000.0) == audit.SEV_MED
    assert audit._severity(0.2, 10.0) == audit.SEV_LOW


def test_findings_are_numbered_by_severity(report):
    order = {audit.SEV_HIGH: 0, audit.SEV_MED: 1, audit.SEV_LOW: 2}
    ranks = [order[f["severity"]] for f in report["findings"]]
    assert ranks == sorted(ranks)
    assert [f["id"] for f in report["findings"]] == \
           [f"F-{i:02d}" for i in range(1, len(report["findings"]) + 1)]


def test_clean_population_gets_a_clean_opinion():
    pop = _pop(rate_periods=[{"id": 1, "valid_from": date(2026, 1, 1),
                              "valid_to": None, "rate_per_km": 3.5,
                              "trip_bonus": 0.0}])
    rep = audit.run_audit(pop)
    assert rep["findings"] == []
    assert rep["opinion"]["level"] == 0


def test_heat_map_covers_every_finding(report):
    assert len(report["heat_map"]) == len(report["findings"])
    for cell in report["heat_map"]:
        assert 1 <= cell["likelihood"] <= 5
        assert 1 <= cell["impact"] <= 5


def test_missing_population_section_is_rejected():
    with pytest.raises(audit.AuditInputError):
        audit.run_audit({"settings": {}})


# ── Выборка ────────────────────────────────────────────────────────────

def test_sample_is_reproducible_across_runs():
    a = audit.sample_indexes(5000, 25, "BEMOL-2026")
    b = audit.sample_indexes(5000, 25, "BEMOL-2026")
    assert a == b and len(a) == 25 and len(set(a)) == 25
    assert a != audit.sample_indexes(5000, 25, "другой ключ")


def test_sample_never_exceeds_the_population():
    assert audit.sample_indexes(10, 25, "k") == list(range(10))


def test_attribute_sample_size_follows_control_frequency():
    assert audit.attribute_sample_size(5000, "daily") == 25
    assert audit.attribute_sample_size(5000, "weekly") == 5
    assert audit.attribute_sample_size(5000, "monthly") == 2
    assert audit.attribute_sample_size(3, "daily") == 3


# ── Книга Excel ────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory, report, population):
    from modules.autopark import audit_excel
    path = str(tmp_path_factory.mktemp("audit") / "book.xlsx")
    return audit_excel.build_workbook(report, population, path)


EXPECTED_SHEETS = ["Титул", "Резюме", "Панель BI", "Контроли", "Находки",
                   "Карта рисков", "Тесты", "Отклонения", "Выборка", "Куб",
                   "КубРейсы", "Самопроверка", "Методика"]


def test_workbook_has_every_sheet(workbook_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    assert wb.sheetnames == EXPECTED_SHEETS


def test_workbook_carries_charts_and_slicers(workbook_path):
    names = zipfile.ZipFile(workbook_path).namelist()
    charts = [n for n in names if n.startswith("xl/charts/chart")]
    assert len(charts) >= 6, f"диаграмм в книге: {len(charts)}"
    from openpyxl import load_workbook
    panel = load_workbook(workbook_path)["Панель BI"]
    assert len(panel.data_validations.dataValidation) >= 4, "нет срезов"


def test_panel_values_are_formulas_not_numbers(workbook_path):
    """Записанное число осталось бы прежним при смене среза."""
    from openpyxl import load_workbook
    panel = load_workbook(workbook_path)["Панель BI"]
    formulas = [c.value for row in panel.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert sum(1 for f in formulas if "SUMPRODUCT" in f) >= 8


def test_formulas_use_comma_as_argument_separator(workbook_path):
    """Точка с запятой — локализованный ввод, в XML файла её быть не может."""
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert ";" not in cell.value, \
                        f"{ws.title}!{cell.coordinate}: {cell.value}"


def test_cube_flag_columns_read_the_panel(workbook_path):
    from openpyxl import load_workbook
    cube = load_workbook(workbook_path)["Куб"]
    flag = cube["H2"].value
    assert isinstance(flag, str) and flag.startswith("=")
    assert "'Панель BI'!$C$6" in flag


def test_demo_workbook_is_marked_as_generated_data(workbook_path):
    from openpyxl import load_workbook
    cover = load_workbook(workbook_path)["Титул"]
    text = " ".join(str(c.value) for row in cover.iter_rows() for c in row
                    if c.value)
    assert "СГЕНЕРИРОВАННЫЙ НАБОР" in text


def test_live_workbook_has_no_selfcheck_sheet(tmp_path):
    """Известного числа дефектов на боевых данных нет — сравнивать не с чем."""
    from modules.autopark import audit_excel
    pop = _pop(rate_periods=[{"id": 1, "valid_from": date(2026, 1, 1),
                              "valid_to": None, "rate_per_km": 3.5}])
    pop["entity"] = "X"
    pop["period"] = {"from": date(2026, 1, 1), "to": date(2026, 9, 30)}
    pop["data_label"] = "БОЕВЫЕ ДАННЫЕ КОНТУРА"
    rep = audit.run_audit(pop)
    path = str(tmp_path / "live.xlsx")
    audit_excel.build_workbook(rep, pop, path)
    from openpyxl import load_workbook
    assert "Самопроверка" not in load_workbook(path).sheetnames


def test_workbook_builds_when_nothing_was_found(tmp_path):
    """Самый правильный исход — пустой реестр — не должен ронять сборку.

    Реестр без строк даёт диапазон условного форматирования «K6:K5», на
    котором openpyxl падает с невнятным TypeError. Проверка «всё чисто»
    обязана доходить до готовой книги.
    """
    from modules.autopark import audit_excel
    pop = _pop(rate_periods=[{"id": 1, "valid_from": date(2026, 1, 1),
                              "valid_to": None, "rate_per_km": 3.5}])
    pop.update({"entity": "X", "data_label": "",
                "period": {"from": date(2026, 1, 1), "to": date(2026, 9, 30)}})
    rep = audit.run_audit(pop)
    assert rep["findings"] == []
    path = str(tmp_path / "clean.xlsx")
    audit_excel.build_workbook(rep, pop, path)
    assert os.path.getsize(path) > 10_000
